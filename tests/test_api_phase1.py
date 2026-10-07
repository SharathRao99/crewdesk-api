import io

import pytest
from openpyxl import load_workbook

KEY = "test-service-key"  # must match conftest.py
ROLE = "php-dev"

ROLE_DATA = {
    "name": ROLE, "languages": ["PHP"], "keywords": ["laravel"], "locations": ["India", "Bengaluru"],
    "exclude_locations": ["Pakistan"], "experience": {"min_years": 2.5, "max_years": 5},
    "primary_skill": {"min_language_share": 0.3, "min_keyword_repos": 1},
    "outreach": {"role_title": "PHP Laravel Developer", "subject": "Role at {company_name}",
                 "body": "Hi {first_name},\nWe like {repo}.\nCareer page: https://example.com/careers"},
}


def cand(login, score, **kw):
    base = {"_id": f"{ROLE}:{login}", "role": ROLE, "login": login, "name": login.title(), "score": score,
            "status": "new", "notes": "", "location": "Bengaluru, India", "est_years": 3.0, "language_share": 0.5,
            "keyword_repos": 3, "language_repos": 5, "matching_repos": [f"{login}-shop"], "last_active": "2026-09-01",
            "url": f"https://github.com/{login}", "public_repos": 10, "followers": 1}
    return {**base, **kw}


@pytest.fixture
def seeded(two_tenants):
    from sourcer import roles_store
    from sourcer.tenancy.db import tenant_db

    for slug in ("acme", "globex"):
        roles_store.import_role(tenant_db(slug), ROLE_DATA)
    col = tenant_db("acme")["candidates"]
    col.delete_many({})
    col.insert_many([
        cand("ann", 90, email="ann@x.com", linkedin_url="https://linkedin.com/in/ann", portfolio_phones=["+919876543210"]),
        cand("bob", 70, portfolio_emails=["bob@y.com"], blog="bob.dev", location="Pune, India"),
        cand("cy", 50, status="contacted", location="Hyderabad, Pakistan"),
        cand("dee", 30, name="Dee", est_years=None),
    ])
    tenant_db("globex")["candidates"].delete_many({})
    tenant_db("globex")["candidates"].insert_one(cand("ann", 60, name="Globex Ann", email="g@z.com"))
    return two_tenants


def H(seeded, tenant="acme"):
    return {"X-Service-Key": KEY, "X-Tenant": tenant, "X-User-Id": seeded[tenant]["user_id"]}


def logins(res):
    return [i["login"] for i in res.json()["items"]]


# ---------- list ----------

def test_list_sorted_and_paginated(api, seeded):
    res = api.get("/candidates", headers=H(seeded))
    assert res.json()["total"] == 4 and logins(res) == ["ann", "bob", "cy", "dee"]
    assert logins(api.get("/candidates?limit=2&skip=1", headers=H(seeded))) == ["bob", "cy"]
    assert logins(api.get("/candidates?sort=name&order=asc", headers=H(seeded))) == ["ann", "bob", "cy", "dee"]


def test_list_item_shape(api, seeded):
    item = api.get("/candidates?q=ann", headers=H(seeded)).json()["items"][0]
    assert item["email"] == "ann@x.com" and item["phone"] == "+919876543210" and item["github_url"] == "https://github.com/ann"
    assert item["linkedin_url"] and item["repos"] == ["ann-shop"] and "bio" not in item


def test_filters(api, seeded):
    get = lambda qs: logins(api.get("/candidates?" + qs, headers=H(seeded)))
    assert get("has_email=true") == ["ann", "bob"]  # bob's email is on his portfolio
    assert get("has_phone=true") == ["ann"]
    assert get("has_linkedin=true") == ["ann"]
    assert get("has_portfolio=true") == ["bob"]
    assert get("status=contacted") == ["cy"]
    assert get("min_score=60") == ["ann", "bob"]
    assert get("location=pune") == ["bob"]
    assert get("has_email=true&min_score=80") == ["ann"]


def test_search_is_literal_not_a_pattern(api, seeded):
    get = lambda q: logins(api.get("/candidates", params={"q": q}, headers=H(seeded)))
    assert get("ANN") == ["ann"]  # case-insensitive
    assert get(".*") == [] and get("c++") == [] and get("(") == []  # regex characters mean nothing


def test_list_rejects_bad_input(api, seeded):
    for qs in ("status=bogus", "sort=password", "limit=1000", "limit=0", "skip=-1", "min_score=500"):
        assert api.get("/candidates?" + qs, headers=H(seeded)).status_code == 422, qs


def test_facets(api, seeded):
    body = api.get("/candidates/facets", headers=H(seeded)).json()
    assert body["total"] == 4
    assert {s["status"]: s["count"] for s in body["by_status"]}["contacted"] == 1


# ---------- detail, patch, events ----------

def test_detail_has_explanation(api, seeded):
    body = api.get(f"/candidates/{ROLE}/ann", headers=H(seeded)).json()
    assert body["emails"] == ["ann@x.com"] and body["why"]["total"] > 0
    assert [g["ok"] for g in body["why"]["gates"]] == [True, True, True]
    assert {r["name"] for r in body["why"]["rules"]} >= {"Framework evidence", "Experience fit"}
    assert api.get(f"/candidates/{ROLE}/nobody", headers=H(seeded)).status_code == 404


def test_detail_without_a_stored_role_still_works(api, seeded):
    from sourcer.tenancy.db import tenant_db

    tenant_db("acme")["roles"].delete_many({})
    body = api.get(f"/candidates/{ROLE}/ann", headers=H(seeded)).json()
    assert body["why"] is None and body["login"] == "ann"


def test_patch_status_and_notes_write_events(api, seeded):
    h = H(seeded)
    res = api.patch(f"/candidates/{ROLE}/bob", json={"status": "contacted", "notes": "emailed Monday"}, headers=h)
    assert res.status_code == 200 and res.json()["status"] == "contacted" and res.json()["notes"] == "emailed Monday"
    events = api.get(f"/candidates/{ROLE}/bob/events", headers=h).json()
    assert {e["type"] for e in events} == {"status", "note"} and all(e["by"] == "hr@acme.com" for e in events)
    status_event = next(e for e in events if e["type"] == "status")
    assert (status_event["from_status"], status_event["to_status"]) == ("new", "contacted")


def test_patch_noop_and_invalid(api, seeded):
    h = H(seeded)
    assert api.patch(f"/candidates/{ROLE}/bob", json={"status": "new"}, headers=h).status_code == 200
    assert api.get(f"/candidates/{ROLE}/bob/events", headers=h).json() == []  # nothing changed, nothing logged
    assert api.patch(f"/candidates/{ROLE}/bob", json={"status": "bogus"}, headers=h).status_code == 422
    assert api.patch(f"/candidates/{ROLE}/bob", json={"notes": "x" * 2001}, headers=h).status_code == 422
    assert api.patch(f"/candidates/{ROLE}/ghost", json={"status": "hired"}, headers=h).status_code == 404


def test_cannot_touch_another_tenants_candidate(api, seeded):
    # Same role and login exist in both tenants: each user only ever reaches their own.
    api.patch(f"/candidates/{ROLE}/ann", json={"status": "hired"}, headers=H(seeded, "globex"))
    assert api.get(f"/candidates/{ROLE}/ann", headers=H(seeded, "acme")).json()["status"] == "new"
    assert api.get(f"/candidates/{ROLE}/ann", headers=H(seeded, "globex")).json()["name"] == "Globex Ann"
    assert api.get(f"/candidates/{ROLE}/bob", headers=H(seeded, "globex")).status_code == 404


# ---------- templates ----------

def tpl(**kw):
    return {"name": "Intro", "subject": "Hi {first_name}", "body": "About {repo}\nhttps://example.com", "role": None, **kw}


def test_template_crud_and_versions(api, seeded):
    h = H(seeded)
    made = api.post("/templates", json=tpl(), headers=h)
    assert made.status_code == 201 and made.json()["version"] == 1
    tid = made.json()["id"]
    updated = api.put(f"/templates/{tid}", json=tpl(subject="Hello {first_name}"), headers=h).json()
    assert updated["version"] == 2 and updated["subject"] == "Hello {first_name}"
    from sourcer.tenancy.db import tenant_db
    stored = tenant_db("acme")["email_templates"].find_one({})
    assert [v["subject"] for v in stored["versions"]] == ["Hi {first_name}"]  # the old text is kept
    assert [t["id"] for t in api.get("/templates", headers=h).json()] == [tid]
    assert api.delete(f"/templates/{tid}", headers=h).status_code == 204
    assert api.get("/templates", headers=h).json() == []


@pytest.mark.parametrize("bad", [
    {"subject": "Hi {firstname}"}, {"body": "x {form_link}"}, {"name": " "}, {"subject": ""}, {"body": "x" * 10001},
])
def test_template_rejects_bad_text(api, seeded, bad):
    assert api.post("/templates", json=tpl(**bad), headers=H(seeded)).status_code == 422


def test_templates_scoped_by_role_and_tenant(api, seeded):
    h = H(seeded)
    api.post("/templates", json=tpl(name="Global"), headers=h)
    api.post("/templates", json=tpl(name="PHP only", role=ROLE), headers=h)
    api.post("/templates", json=tpl(name="Other role", role="java-dev"), headers=h)
    assert [t["name"] for t in api.get(f"/templates?role={ROLE}", headers=h).json()] == ["Global", "PHP only"]
    assert api.get("/templates", headers=H(seeded, "globex")).json() == []
    other = api.get("/templates", headers=h).json()[0]["id"]
    assert api.put(f"/templates/{other}", json=tpl(), headers=H(seeded, "globex")).status_code == 404
    assert api.delete(f"/templates/{other}", headers=H(seeded, "globex")).status_code == 404


def test_preview_flags_unknown_and_fills_values(api, seeded):
    h = H(seeded)
    sample = api.post("/templates/preview", json={"subject": "{company_name}", "body": "{first_name} {nope}"}, headers=h).json()
    assert sample["sample"] is True and sample["subject"] == "Acme Inc" and sample["unknown"] == ["nope"]
    real = api.post("/templates/preview", json={"subject": "s", "body": "{first_name} / {repo} / {role_title}", "role": ROLE, "login": "ann"}, headers=h).json()
    assert real["sample"] is False and real["body"] == "Ann / ann-shop / PHP Laravel Developer"


def test_placeholder_list(api, seeded):
    names = {p["name"] for p in api.get("/templates/placeholders", headers=H(seeded)).json()}
    assert {"first_name", "repo", "company_name", "role_title", "sender"} <= names


# ---------- drafts and roles ----------

def test_draft_uses_role_outreach_by_default(api, seeded):
    body = api.post(f"/candidates/{ROLE}/ann/draft", json={"sender": "Sam"}, headers=H(seeded)).json()
    assert body["to"] == "ann@x.com" and body["subject"] == "Role at Acme Inc"
    assert body["body"].endswith("https://example.com/careers")
    assert body["gmail_link"].startswith("https://mail.google.com/") and body["mailto_link"].startswith("mailto:ann@x.com?")
    assert body["mailto_link"].endswith("%0D%0A%0D%0A")  # trailing breaks so the last URL becomes a link


def test_draft_with_template_and_role_default(api, seeded):
    h = H(seeded)
    tid = api.post("/templates", json=tpl(subject="From {sender}", body="Hello {full_name}"), headers=h).json()["id"]
    assert api.post(f"/candidates/{ROLE}/ann/draft", json={"template_id": tid, "sender": "Sam"}, headers=h).json()["subject"] == "From Sam"
    assert api.put(f"/roles/{ROLE}/default-template", json={"template_id": tid}, headers=h).json()["default_template_id"] == tid
    assert api.post(f"/candidates/{ROLE}/ann/draft", json={}, headers=h).json()["body"] == "Hello Ann"
    api.delete(f"/templates/{tid}", headers=h)  # deleting a template clears it as the role default
    assert api.get("/roles", headers=h).json()[0]["default_template_id"] is None


def test_draft_edge_cases(api, seeded):
    h = H(seeded)
    nobody_email = api.post(f"/candidates/{ROLE}/dee/draft", json={}, headers=h).json()
    assert nobody_email["to"] is None and nobody_email["gmail_link"] is None
    assert api.post(f"/candidates/{ROLE}/ann/draft", json={"template_id": "0" * 24}, headers=h).status_code == 404
    # another tenant's template cannot be used
    tid = api.post("/templates", json=tpl(), headers=H(seeded, "globex")).json()["id"]
    assert api.post(f"/candidates/{ROLE}/ann/draft", json={"template_id": tid}, headers=h).status_code == 404


def test_roles_list_counts(api, seeded):
    roles = api.get("/roles", headers=H(seeded)).json()
    assert roles[0]["name"] == ROLE and roles[0]["candidates"] == 4 and roles[0]["keywords"] == ["laravel"]
    assert api.put("/roles/ghost/default-template", json={"template_id": None}, headers=H(seeded)).status_code == 404


# ---------- export ----------

def test_export_csv_contact_only_with_drafts_and_audit(api, seeded):
    import csv

    res = api.get(f"/export?role={ROLE}&format=csv&sender=Sam", headers=H(seeded))
    assert res.status_code == 200 and res.headers["content-disposition"] == f'attachment; filename="{ROLE}.csv"'
    rows = list(csv.DictReader(io.StringIO(res.content.decode("utf-8-sig"))))
    # cy is in Pakistan (excluded) and dee has no contact; both are left out.
    assert [r["login"] for r in rows] == ["ann", "bob"]
    assert rows[0]["email_to"] == "ann@x.com" and rows[0]["email_subject"] == "Role at Acme Inc"
    assert rows[1]["email_to"] == "bob@y.com" and rows[1]["gmail_link"].startswith("https://mail.google.com/")
    from sourcer.tenancy.db import tenant_db
    log = tenant_db("acme")["audit_log"].find_one({"action": "export"})
    assert log["email"] == "hr@acme.com" and log["detail"]["rows"] == 2


def test_export_xlsx_has_clickable_links_and_is_isolated(api, seeded):
    res = api.get(f"/export?role={ROLE}", headers=H(seeded))
    ws = load_workbook(io.BytesIO(res.content))["Candidates"]
    headers = [c.value for c in ws[1]]
    link_cell = ws.cell(row=2, column=headers.index("gmail_link") + 1)
    assert link_cell.value == "Open in Gmail" and link_cell.hyperlink.target.startswith("https://mail.google.com/")
    globex = api.get(f"/export?role={ROLE}&format=csv", headers=H(seeded, "globex")).content.decode("utf-8-sig")
    assert "Globex Ann" in globex and "ann@x.com" not in globex and "Bob" not in globex


def test_export_filters_and_validation(api, seeded):
    import csv

    res = api.get(f"/export?role={ROLE}&format=csv&require=none&status=new", headers=H(seeded))
    rows = list(csv.DictReader(io.StringIO(res.content.decode("utf-8-sig"))))
    # require=none keeps dee (no contact); cy is excluded by status and by location.
    assert [r["login"] for r in rows] == ["ann", "bob", "dee"]
    assert api.get("/export", headers=H(seeded)).status_code == 422  # role is required
    assert api.get(f"/export?role={ROLE}&format=pdf", headers=H(seeded)).status_code == 422
    assert api.get(f"/export?role={ROLE}&require=everything", headers=H(seeded)).status_code == 422


# ---------- account ----------

def test_me_and_password_change(api, seeded):
    h = H(seeded)
    me = api.get("/me", headers=h).json()
    assert me["email"] == "hr@acme.com" and me["tenant_name"] == "Acme Inc" and me["must_change_password"] is True
    assert api.post("/me/password", json={"old": "wrong", "new": "a-long-enough-password"}, headers=h).status_code == 400
    assert api.post("/me/password", json={"old": seeded["acme"]["password"], "new": "short"}, headers=h).status_code == 422
    assert api.post("/me/password", json={"old": seeded["acme"]["password"], "new": "a-long-enough-password"}, headers=h).status_code == 204
    assert api.get("/me", headers=h).json()["must_change_password"] is False
    login = api.post("/auth/login", json={"email": "hr@acme.com", "password": "a-long-enough-password"}, headers={"X-Service-Key": KEY})
    assert login.status_code == 200


# ---------- admin commands ----------

def test_revalidate_removes_only_uncontacted_excluded_candidates(seeded):
    import argparse

    from sourcer import admin
    from sourcer.tenancy.db import tenant_db

    col = tenant_db("acme")["candidates"]
    col.insert_one(cand("eve", 10, location="Karachi, Pakistan"))
    admin._candidates_revalidate(argparse.Namespace(tenant="acme", role=ROLE))
    left = sorted(d["login"] for d in col.find())
    assert "eve" not in left       # new + excluded: removed
    assert "cy" in left            # already contacted: kept for a human to review
