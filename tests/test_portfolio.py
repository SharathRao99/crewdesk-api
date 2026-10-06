from sourcer.portfolio import normalize_url, parse_page

HTML = """
<a href="https://www.linkedin.com/in/jane-dev/?trk=x">LinkedIn</a>
<a href="mailto:jane@example.com?subject=hi">Email me</a>
<a href="/files/jane.pdf">Download CV</a>
<a href="/about">About</a><a href="/blog/post">Blog</a>
<img src="logo@2x.png"> Reach me: hello@jane.dev
"""


def test_parse_page_finds_contact_links():
    found = parse_page(HTML, "https://jane.dev/")
    assert found["linkedin"] == {"https://www.linkedin.com/in/jane-dev"}
    assert found["emails"] == {"jane@example.com", "hello@jane.dev"}
    assert found["resumes"] == {"https://jane.dev/files/jane.pdf"}
    assert found["pages"] == {"https://jane.dev/about"}


def test_normalize_url():
    assert normalize_url("jane.dev") == "https://jane.dev"
    assert normalize_url("") is None
    assert normalize_url("http://localhost:3000") is None
    assert normalize_url("http://192.168.1.5") is None
    assert normalize_url("https://github.com/jane") is None


def test_phone_extraction():
    from sourcer.portfolio import find_phones

    html = '<a href="tel:+91 98765 43210">Call</a> or +44 7700 900123. Order 1234567890 shipped.'
    found = parse_page(html, "https://jane.dev/")
    assert found["phones"] == {"+919876543210", "+447700900123"}
    assert find_phones("no number here 2025-10-05") == set()


def test_report_embeds_safe_json(tmp_path):
    from sourcer.report import build_report
    from sourcer.role import Role

    role = Role(name="r", languages=["PHP"], keywords=[], locations=["India"])
    row = {"login": "a", "name": "</script><b>x", "url": "u", "score": 50, "linkedin_url": "https://linkedin.com/in/a"}
    out = tmp_path / "report.html"
    assert build_report(role, [row, {"login": "b", "url": "u", "score": 1}], str(out)) == 1  # b has no useful link
    html = out.read_text()
    assert "</script><b>x" not in html and "<\\/script>" in html


def test_csv_includes_email_draft_and_skips_rows_without_links(tmp_path):
    import csv

    from sourcer.export import write_csv

    outreach = {"role_title": "Laravel Developer", "subject": "Role for {first_name}", "body": "Hi {first_name}, loved {repo}.\n{sender}"}
    rows = [
        {"login": "a", "name": "Jane Doe", "email": "j@x.com", "linkedin_url": "https://linkedin.com/in/j", "matching_repos": ["shop"], "score": 80},
        {"login": "b", "name": "No Links", "email": "n@x.com", "score": 90},
    ]
    out = tmp_path / "c.csv"
    assert write_csv(str(out), rows, outreach, "Sam") == 1
    row = next(csv.DictReader(out.open()))
    assert row["email_to"] == "j@x.com"
    assert row["email_subject"] == "Role for Jane"
    assert row["email_message"] == "Hi Jane, loved shop.\nSam"


def test_csv_contact_filter_keeps_email_or_phone(tmp_path):
    import csv

    from sourcer.export import has_contact, write_csv

    rows = [
        {"login": "e", "email": "e@x.com", "score": 1},
        {"login": "p", "portfolio_phones": ["+919876543210"], "score": 2},
        {"login": "m", "portfolio_emails": ["m@x.com"], "score": 3},
        {"login": "n", "linkedin_url": "https://linkedin.com/in/n", "score": 4},
    ]
    out = tmp_path / "c.csv"
    assert write_csv(str(out), rows, keep=has_contact) == 3
    assert [r["login"] for r in csv.DictReader(out.open())] == ["e", "p", "m"]


def test_mail_links_round_trip_subject_and_body():
    from urllib.parse import parse_qs, urlparse

    from sourcer.outreach import mail_links

    body = "Hi Jane,\n\nVisit https://x.com/?a=1&b=2 now!\nRegards"
    links = mail_links("j@x.com", "Role & more", body)
    for key in ("mailto_link", "gmail_link", "outlook_link"):
        parsed = urlparse(links[key])
        query = parse_qs(parsed.query)
        assert query.get("subject", query.get("su"))[0] == "Role & more"
        got = query["body"][0]
        assert got.replace("\r\n", "\n") == body + "\n\n"  # trailing breaks so a final URL is linkified
    assert links["mailto_link"].startswith("mailto:j@x.com?")
    assert "%0D%0A" in links["mailto_link"]  # CRLF line breaks for mail apps


def test_xlsx_has_clickable_mail_links(tmp_path):
    from openpyxl import load_workbook

    from sourcer.export import COLUMNS, has_contact, write_xlsx

    outreach = {"role_title": "Dev", "subject": "Hi {first_name}", "body": "Hello {first_name}"}
    out = tmp_path / "c.xlsx"
    assert write_xlsx(str(out), [{"login": "a", "name": "Jane Doe", "email": "j@x.com", "score": 5}], outreach, keep=has_contact) == 1
    cell = load_workbook(out)["Candidates"].cell(row=2, column=COLUMNS.index("gmail_link") + 1)
    assert cell.value == "Open in Gmail" and cell.hyperlink.target.startswith("https://mail.google.com/")
