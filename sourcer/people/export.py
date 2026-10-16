"""Leave and balances as CSV or Excel files. Cells that start with a formula character are defused, because reasons
and notes are text people typed and a spreadsheet would otherwise run them."""

import csv
import io
from typing import Any, Dict, List, Sequence, Tuple

from pymongo.database import Database

from sourcer.people import leave, policy as pol, store

CSV_MEDIA = "text/csv; charset=utf-8"
XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")

LEAVE_HEADERS = ["Employee", "Employee no.", "Leave type", "From", "To", "Day part", "Days", "Status", "Over balance", "Reason", "Note", "Approved by", "Recorded on"]
BALANCE_FIXED = ["Employee", "Employee no.", "Employee status", "Leave year"]
_PARTS = {"full": "Full day", "first_half": "First half", "second_half": "Second half"}


def safe(value: Any) -> Any:
    """Text that a spreadsheet could take for a formula gets a leading apostrophe; numbers and empty cells pass through."""
    if isinstance(value, str) and value.startswith(_FORMULA_START):
        return "'" + value
    return value


def csv_file(headers: Sequence[str], rows: List[Sequence[Any]]) -> bytes:
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer)
    writer.writerow(headers)
    for row in rows:
        writer.writerow(["" if v is None else safe(v) for v in row])
    return buffer.getvalue().encode("utf-8-sig")  # BOM so Excel reads UTF-8 names correctly


def xlsx_file(sheet: str, headers: Sequence[str], rows: List[Sequence[Any]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = sheet[:31]
    ws.append(list(headers))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for row in rows:
        ws.append(["" if v is None else safe(v) for v in row])
    for i, h in enumerate(headers, start=1):
        longest = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:200] if r[i - 1] is not None])
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(max(12, longest + 2), 48)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def render(fmt: str, sheet: str, headers: Sequence[str], rows: List[Sequence[Any]]):
    """(bytes, media type) for "csv" or "xlsx"."""
    return (csv_file(headers, rows), CSV_MEDIA) if fmt == "csv" else (xlsx_file(sheet, headers, rows), XLSX_MEDIA)


def _people(db: Database, ids: List[Any]) -> Dict[str, Dict[str, Any]]:
    wanted = list({str(i) for i in ids if i})
    docs = db["employees"].find({"_id": {"$in": [store.oid(i) for i in wanted]}}, {"name": 1, "employee_no": 1}) if wanted else []
    return {str(d["_id"]): d for d in docs}


def leave_rows(db: Database, records: List[Dict[str, Any]]) -> List[List[Any]]:
    people = _people(db, [r["employee_id"] for r in records])
    rows = []
    for r in records:
        who = people.get(str(r["employee_id"]), {})
        created = r.get("created_at")
        rows.append([
            who.get("name", ""), who.get("employee_no"), r["type_name"], r["start_date"], r["end_date"], _PARTS.get(r["day_part"], r["day_part"]),
            r["days"], str(r["status"]).capitalize(), "Yes" if r.get("over_balance") else "No", r.get("reason"), r.get("note"),
            r.get("approved_by_name"), created.date().isoformat() if created else None,
        ])
    return rows


def balance_table(db: Database, year: int = None) -> Tuple[List[str], List[List[Any]]]:
    """One row per employee, with a "used" and a "remaining" column for each leave type, for a leave year (default: the current one)."""
    people = list(db["employees"].find({}, {"name": 1, "employee_no": 1, "joining_date": 1, "status": 1}).sort("name_lc", 1).limit(2000))
    policy = pol.get_policy(db)
    order: List[str] = []  # leave types in the company's own order, only those that appear for someone
    per_person = []
    for emp in people:
        b = leave.balances(db, emp, year, policy)
        per_person.append((emp, b))
        for t in b["types"]:
            if t["name"] not in order:
                order.append(t["name"])
    headers = list(BALANCE_FIXED)
    for name in order:
        headers += [f"{name} used", f"{name} remaining"]
    rows = []
    for emp, b in per_person:
        by_name = {t["name"]: t for t in b["types"]}
        row: List[Any] = [emp["name"], emp.get("employee_no"), str(emp.get("status", "")).capitalize(), b["label"]]
        for name in order:
            t = by_name.get(name)
            row += [None, None] if t is None else [t["used"], "No limit" if t["remaining"] is None else t["remaining"]]
        rows.append(row)
    return headers, rows
