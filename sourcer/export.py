"""CSV export for manual outreach."""

import csv
import io
from typing import Any, Callable, Dict, Iterable, List, Optional

from sourcer.outreach import best_email, mail_links, render_draft

COLUMNS = [
    "score", "name", "email_to", "gmail_link", "outlook_link", "mailto_link", "email_subject", "email_message",
    "linkedin_url", "portfolio_url", "resume_links", "portfolio_phones", "location", "company", "est_years",
    "last_active", "matching_repos", "hireable", "bio", "url", "login", "status", "notes",
]
LINK_LABELS = {"gmail_link": "Open in Gmail", "outlook_link": "Open in Outlook", "mailto_link": "Open in mail app"}


def flatten(row: Dict[str, Any]) -> Dict[str, Any]:
    """Join list fields into comma-separated strings for CSV."""
    return {
        **row,
        **{k: ", ".join(row.get(k) or []) for k in ("matching_repos", "portfolio_emails", "resume_links", "portfolio_phones")},
    }


def has_useful_link(row: Dict[str, Any]) -> bool:
    """Only LinkedIn, portfolio and resume links count; other socials are ignored."""
    return bool(row.get("linkedin_url") or row.get("portfolio_url") or row.get("blog") or row.get("resume_links"))


def has_contact(row: Dict[str, Any]) -> bool:
    """True if we can reach the person directly: any email (GitHub or portfolio) or a phone number."""
    return bool(best_email(row) or row.get("portfolio_phones"))


def prepare(
    rows: Iterable[Dict[str, Any]],
    outreach: Optional[Dict[str, str]],
    sender: str,
    keep: Callable[[Dict[str, Any]], bool],
    company_name: str = "",
) -> List[Dict[str, Any]]:
    """Filter rows with `keep` and add the email address, draft and mail links."""
    out = []
    for row in rows:
        if not keep(row):
            continue
        row = {**row, "email_to": best_email(row)}
        if outreach and row["email_to"]:
            row["email_subject"], row["email_message"] = render_draft(row, outreach, sender, company_name)
            row.update(mail_links(row["email_to"], row["email_subject"], row["email_message"]))
        out.append(flatten(row))
    return out


def csv_bytes(prepared: List[Dict[str, Any]]) -> bytes:
    buffer = io.StringIO(newline="")
    writer = csv.DictWriter(buffer, fieldnames=COLUMNS, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(prepared)
    return buffer.getvalue().encode("utf-8-sig")  # BOM so Excel reads UTF-8 names correctly


def write_csv(
    path: str,
    rows: Iterable[Dict[str, Any]],
    outreach: Optional[Dict[str, str]] = None,
    sender: str = "",
    keep: Callable[[Dict[str, Any]], bool] = has_useful_link,
    company_name: str = "",
) -> int:
    """Write rows that pass `keep` (default: has a LinkedIn/portfolio/resume link). Returns the number written.

    With `outreach`, rows that have an email get a ready-to-paste draft and mail links.
    """
    prepared = prepare(rows, outreach, sender, keep, company_name)
    with open(path, "wb") as f:
        f.write(csv_bytes(prepared))
    return len(prepared)


def xlsx_bytes(prepared: List[Dict[str, Any]]) -> bytes:
    """The same data as the CSV, but the mail links are real clickable cells."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Candidates"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"

    for r, row in enumerate(prepared, start=2):
        for c, column in enumerate(COLUMNS, start=1):
            value = row.get(column)
            if value is None or value == "":
                continue
            cell = ws.cell(row=r, column=c)
            if column in LINK_LABELS or column in ("linkedin_url", "portfolio_url"):
                cell.value = LINK_LABELS.get(column, value)
                cell.hyperlink = value
                cell.font = Font(color="0563C1", underline="single")
            else:
                cell.value = value
    widths = {"name": 24, "email_to": 30, "gmail_link": 16, "outlook_link": 18, "mailto_link": 18, "email_subject": 36, "email_message": 50}
    for c, column in enumerate(COLUMNS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = widths.get(column, 18)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def write_xlsx(
    path: str,
    rows: Iterable[Dict[str, Any]],
    outreach: Optional[Dict[str, str]] = None,
    sender: str = "",
    keep: Callable[[Dict[str, Any]], bool] = has_useful_link,
    company_name: str = "",
) -> int:
    prepared = prepare(rows, outreach, sender, keep, company_name)
    with open(path, "wb") as f:
        f.write(xlsx_bytes(prepared))
    return len(prepared)
